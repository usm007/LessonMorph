"""DOCX document extractor using python-docx."""

from __future__ import annotations
from pathlib import Path
from typing import List, Optional
import docx
from lessonmorph.core.models import DocumentMap, DocumentSection
from lessonmorph.ingest.base import BaseExtractor, ExtractedImage, ExtractedTable, IngestionResult


class DocxExtractor(BaseExtractor):
    """Extractor for Microsoft Word (.docx) documents."""

    def extract(self, file_path: Path) -> IngestionResult:
        doc = docx.Document(file_path)
        title = file_path.stem.replace("_", " ").title()

        sections: List[DocumentSection] = []
        tables: List[ExtractedTable] = []
        images: List[ExtractedImage] = []

        # Parse tables
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])
            if table_rows:
                headers = table_rows[0]
                body_rows = table_rows[1:] if len(table_rows) > 1 else []
                tables.append(
                    ExtractedTable(
                        title=f"Table {t_idx + 1}",
                        headers=headers,
                        rows=body_rows,
                        page_number=1,
                        source_ref=f"DOCX Table {t_idx + 1}",
                    )
                )

        # Parse paragraphs and detect headings
        current_sec_id = 1
        current_sec_title = "Introduction"
        current_sec_text: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            if "Heading 1" in style_name or "Title" in style_name or (style_name == "Heading" and len(text) < 60):
                if current_sec_text:
                    sections.append(
                        DocumentSection(
                            id=f"ch{current_sec_id:02d}",
                            title=current_sec_title,
                            start_page=current_sec_id,
                            end_page=current_sec_id,
                            text_content="\n\n".join(current_sec_text),
                        )
                    )
                    current_sec_id += 1
                    current_sec_text = []
                current_sec_title = text
            else:
                current_sec_text.append(text)

        if current_sec_text or not sections:
            sections.append(
                DocumentSection(
                    id=f"ch{current_sec_id:02d}",
                    title=current_sec_title,
                    start_page=current_sec_id,
                    end_page=current_sec_id,
                    text_content="\n\n".join(current_sec_text),
                )
            )

        doc_map = DocumentMap(
            title=title,
            source_file=str(file_path),
            total_pages=len(sections),
            sections=sections,
            metadata={"format": "docx"},
        )

        return IngestionResult(
            document_title=title,
            source_path=file_path,
            total_pages=len(sections),
            doc_map=doc_map,
            tables=tables,
            images=images,
        )
