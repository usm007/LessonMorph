"""Shape and component helpers for PowerPoint rendering."""

from __future__ import annotations
from typing import Optional
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from lessonmorph.renderer.design import DesignSystem


class ShapeHelper:
    """Helper methods for constructing consistent, styled PowerPoint shapes."""

    @classmethod
    def set_slide_background(cls, slide, color: RGBColor = DesignSystem.COLOR_BG_PAGE) -> None:
        """Sets the slide background color using a full-bleed background shape."""
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            DesignSystem.SLIDE_WIDTH,
            DesignSystem.SLIDE_HEIGHT,
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()  # no line
        # Send background to back by reordering XML
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)

    @classmethod
    def add_header(
        cls,
        slide,
        title: str,
        subtitle: Optional[str] = None,
        category: Optional[str] = None,
    ) -> None:
        """Standardized, high-contrast slide header with category pill."""
        top_offset = DesignSystem.MARGIN_TOP

        # 1. Category Badge if present
        if category:
            badge = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                DesignSystem.MARGIN_LEFT,
                top_offset,
                Inches(2.5),
                Inches(0.32),
            )
            badge.fill.solid()
            badge.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT_BG
            badge.line.color.rgb = DesignSystem.COLOR_ACCENT_BORDER
            badge.line.width = Pt(1)
            tf = badge.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = category.upper()
            p.font.name = DesignSystem.FONT_BODY
            p.font.size = DesignSystem.SIZE_SECTION_BADGE
            p.font.bold = True
            p.font.color.rgb = DesignSystem.COLOR_ACCENT
            p.alignment = PP_ALIGN.CENTER
            top_offset += Inches(0.38)

        # 2. Main Title Text Box
        title_box = slide.shapes.add_textbox(
            DesignSystem.MARGIN_LEFT,
            top_offset,
            DesignSystem.CONTENT_WIDTH,
            Inches(0.65),
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = DesignSystem.FONT_TITLE
        p.font.size = DesignSystem.SIZE_SLIDE_TITLE
        p.font.bold = True
        p.font.color.rgb = DesignSystem.COLOR_TEXT_PRIMARY

        # 3. Subtitle
        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = DesignSystem.FONT_BODY
            p2.font.size = DesignSystem.SIZE_SUBTITLE
            p2.font.color.rgb = DesignSystem.COLOR_TEXT_SECONDARY
            p2.space_before = Pt(4)

    @classmethod
    def add_card(
        cls,
        slide,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
        bg_color: RGBColor = DesignSystem.COLOR_BG_CARD,
        border_color: RGBColor = DesignSystem.COLOR_BORDER_CARD,
        border_width: Pt = Pt(1.5),
    ):
        """Creates a clean rounded-rectangle card container."""
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left,
            top,
            width,
            height,
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = border_width
        return card
