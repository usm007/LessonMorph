"""Semantic renderers: visual composition for the Presentation IR grammar.

Each renderer executes ONE Representation from structured elements_data.
Pure execution: no pedagogy, no reinterpretation, no invented content.
Real PowerPoint objects — shapes, connectors, arrows, native tables —
never prose dumped into decorative cards.
"""

from __future__ import annotations
from typing import Any, Dict, List, Tuple
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from lessonmorph.core.models import AnimationType, SlideSpec
from lessonmorph.renderer.design import DesignSystem
from lessonmorph.renderer.shapes import ShapeHelper


def _textbox(slide, left, top, width, height, runs, size=Pt(15),
             bold=False, color=None, align=None, name="Segoe UI") -> None:
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for text, opts in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = text
        p.font.name = name
        p.font.size = opts.get("size", size)
        p.font.bold = opts.get("bold", bold)
        p.font.color.rgb = opts.get("color", color or DesignSystem.COLOR_TEXT_PRIMARY)
        if opts.get("align", align) is not None:
            p.alignment = opts.get("align", align)
        if opts.get("space_before"):
            p.space_before = opts["space_before"]


def _badge(slide, left, top, width, text, bg, fg, size=Pt(11)) -> None:
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.34))
    b.fill.solid()
    b.fill.fore_color.rgb = bg
    b.line.fill.background()
    p = b.text_frame.paragraphs[0]
    p.text = text.upper()
    p.font.name = DesignSystem.FONT_BODY
    p.font.size = size
    p.font.bold = True
    p.font.color.rgb = fg
    p.alignment = PP_ALIGN.CENTER


def _connector(slide, x1, y1, x2, y2, color=None, width=Pt(2)) -> None:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color or DesignSystem.COLOR_ACCENT
    c.line.width = width


class SemanticRenderer:
    """Visual composer for IR representations."""

    # -- structural ------------------------------------------------------
    @classmethod
    def render_labeled_diagram(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Labeled Diagram")
        subject, labels = d.get("subject", spec.title), d.get("labels", [])[:6]
        cx, cy, cw, ch = Inches(5.4), Inches(2.6), Inches(2.6), Inches(2.6)
        hub = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
        hub.fill.solid()
        hub.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT_BG
        hub.line.color.rgb = DesignSystem.COLOR_ACCENT
        hub.line.width = Pt(2.5)
        p = hub.text_frame.paragraphs[0]
        p.text = subject[:80]
        p.font.name = DesignSystem.FONT_TITLE
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DesignSystem.COLOR_TEXT_PRIMARY
        p.alignment = PP_ALIGN.CENTER
        hub.text_frame.word_wrap = True
        animated = [(hub.shape_id, AnimationType.APPEAR)]
        if not labels:
            _textbox(slide, Inches(1.2), Inches(5.5), Inches(10.9), Inches(1.2),
                     [(d.get("caption", "")[:300], {})], size=Pt(14),
                     color=DesignSystem.COLOR_TEXT_SECONDARY)
            return animated
        left = labels[::2] or labels[:3]
        right = labels[1::2] or []
        col_w = Inches(3.4)
        for col, items in ((DesignSystem.MARGIN_LEFT, left),
                           (Inches(13.333 - 0.8 - 3.4), right)):
            for i, lab in enumerate(items):
                y = Inches(2.3) + i * Inches(1.55)
                card = ShapeHelper.add_card(slide, col, y, col_w, Inches(1.35),
                                            border_color=DesignSystem.COLOR_ACCENT_BORDER)
                _textbox(slide, col + Inches(0.2), y + Inches(0.08),
                         col_w - Inches(0.4), Inches(1.2),
                         [(lab.get("name", "")[:80], {"bold": True, "size": Pt(13),
                                                      "color": DesignSystem.COLOR_ACCENT}),
                          (lab.get("explanation", "")[:140],
                           {"size": Pt(12), "space_before": Pt(4)})])
                x1 = col + (col_w if col < Inches(6) else Inches(0))
                _connector(slide, x1, y + Inches(0.67),
                           cx + (Inches(0) if col < Inches(6) else cw),
                           cy + Inches(1.3))
                animated.append((card.shape_id, AnimationType.APPEAR))
        return animated

    # -- process ----------------------------------------------------------
    @classmethod
    def render_chain(cls, slide, spec: SlideSpec, cycle: bool = False) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle,
                               "Cycle" if cycle else "Pathway")
        stages = (d.get("stages") or [])[:6]
        if not stages:
            stages = (d.get("points") or ["Stage from source."])[:4]
        n = len(stages)
        animated: List[Tuple[int, AnimationType]] = []
        if cycle and n >= 3:
            cx, cy, rx, ry = Inches(6.67), Inches(4.4), Inches(3.6), Inches(1.5)
            import math
            pos = []
            for i in range(n):
                a = -math.pi / 2 + 2 * math.pi * i / n
                pos.append((cx + rx * math.cos(a) - Inches(1.35),
                            cy + ry * math.sin(a) - Inches(0.55)))
            for i, (x, y) in enumerate(pos):
                card = ShapeHelper.add_card(slide, x, y, Inches(2.7), Inches(1.1),
                                            border_color=DesignSystem.COLOR_ACCENT)
                _textbox(slide, x + Inches(0.2), y + Inches(0.12),
                         Inches(2.3), Inches(0.9),
                         [(f"STEP {i+1}", {"bold": True, "size": Pt(11),
                                           "color": DesignSystem.COLOR_ACCENT}),
                          (stages[i][:140], {"size": Pt(12), "space_before": Pt(2)})])
                animated.append((card.shape_id, AnimationType.APPEAR))
            for i in range(n):
                x1, y1 = pos[i][0] + Inches(1.35), pos[i][1] + Inches(0.55)
                x2, y2 = pos[(i + 1) % n][0] + Inches(1.35), pos[(i + 1) % n][1] + Inches(0.55)
                _connector(slide, x1, y1, x2, y2)
            return animated
        gap, top_y, h = Inches(0.75), Inches(2.6), Inches(3.2)
        node_w = (DesignSystem.CONTENT_WIDTH - gap * n) / n
        for i, st in enumerate(stages):
            x = DesignSystem.MARGIN_LEFT + i * (node_w + gap)
            card = ShapeHelper.add_card(slide, x, top_y, node_w, h,
                                        border_color=DesignSystem.COLOR_ACCENT
                                        if i == 0 else DesignSystem.COLOR_BORDER_CARD,
                                        border_width=Pt(2) if i == 0 else Pt(1.5))
            _textbox(slide, x + Inches(0.2), top_y + Inches(0.2), node_w - Inches(0.4), h - Inches(0.4),
                     [(f"STEP {i+1}", {"bold": True, "size": Pt(12),
                                       "color": DesignSystem.COLOR_ACCENT}),
                      (st[:260], {"size": Pt(13), "space_before": Pt(6)})])
            animated.append((card.shape_id, AnimationType.APPEAR))
            if i < n - 1:
                ax = x + node_w
                arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax - Inches(0.02),
                                             top_y + h / 2 - Inches(0.18),
                                             gap + Inches(0.04), Inches(0.36))
                arr.fill.solid()
                arr.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT
                arr.line.fill.background()
        return animated

    # -- quantitative ------------------------------------------------------
    @classmethod
    def render_equation(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        eq = d.get("equation", {})
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Equation")
        inline = eq.get("inline", "")
        eq_card = ShapeHelper.add_card(slide, DesignSystem.MARGIN_LEFT, Inches(2.1),
                                       DesignSystem.CONTENT_WIDTH, Inches(1.5),
                                       bg_color=DesignSystem.COLOR_BG_HERO,
                                       border_color=DesignSystem.COLOR_BG_HERO)
        _textbox(slide, DesignSystem.MARGIN_LEFT + Inches(0.4), Inches(2.35),
                 DesignSystem.CONTENT_WIDTH - Inches(0.8), Inches(1.1),
                 [(inline[:220], {"size": Pt(26), "bold": True,
                                  "color": DesignSystem.COLOR_TEXT_ON_DARK,
                                  "align": PP_ALIGN.CENTER})], name="Consolas")
        terms = [t for t in (eq.get("lhs", []) + eq.get("rhs", [])) if t.get("note")]
        animated = [(eq_card.shape_id, AnimationType.APPEAR)]
        for i, t in enumerate(terms[:4]):
            col, row = i % 2, i // 2
            cw = (DesignSystem.CONTENT_WIDTH - Inches(0.4)) / 2
            x = DesignSystem.MARGIN_LEFT + col * (cw + Inches(0.4))
            y = Inches(3.9) + row * Inches(1.5)
            card = ShapeHelper.add_card(slide, x, y, cw, Inches(1.3))
            _textbox(slide, x + Inches(0.25), y + Inches(0.12), cw - Inches(0.5), Inches(1.1),
                     [(f"{t.get('coefficient','')} {t.get('species','')}".strip(),
                       {"bold": True, "size": Pt(15)}),
                      (t.get("note", "")[:160], {"size": Pt(12), "space_before": Pt(4)})])
            animated.append((card.shape_id, AnimationType.APPEAR))
        if d.get("context"):
            _textbox(slide, DesignSystem.MARGIN_LEFT, Inches(6.7),
                     DesignSystem.CONTENT_WIDTH, Inches(0.5),
                     [(d["context"][:200], {"size": Pt(12)})],
                     color=DesignSystem.COLOR_TEXT_SECONDARY)
        return animated

    @classmethod
    def render_worked_calc(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Worked Calculation")
        y = Inches(2.0)
        animated: List[Tuple[int, AnimationType]] = []
        givens = d.get("givens", [])
        if givens:
            _badge(slide, DesignSystem.MARGIN_LEFT, y, Inches(1.6),
                   "Givens", DesignSystem.COLOR_ACCENT_BG, DesignSystem.COLOR_ACCENT)
            _textbox(slide, DesignSystem.MARGIN_LEFT + Inches(1.8), y - Inches(0.04),
                     DesignSystem.CONTENT_WIDTH - Inches(1.8), Inches(0.45),
                     [("  •  ".join(givens)[:260], {})], size=Pt(13),
                     color=DesignSystem.COLOR_TEXT_SECONDARY)
            y += Inches(0.5)
        _textbox(slide, DesignSystem.MARGIN_LEFT, y, DesignSystem.CONTENT_WIDTH, Inches(0.9),
                 [(d.get("problem", "")[:320], {"bold": True})], size=Pt(15))
        y += Inches(0.95)
        steps = d.get("steps", [])[:4]
        n = max(1, len(steps))
        sw = (DesignSystem.CONTENT_WIDTH - Inches(0.3 * (n - 1))) / n
        for i, st in enumerate(steps):
            x = DesignSystem.MARGIN_LEFT + i * (sw + Inches(0.3))
            card = ShapeHelper.add_card(slide, x, y, sw, Inches(2.5),
                                        border_color=DesignSystem.COLOR_ACCENT
                                        if i == len(steps) - 1 else DesignSystem.COLOR_BORDER_CARD,
                                        border_width=Pt(2) if i == len(steps) - 1 else Pt(1.5))
            _textbox(slide, x + Inches(0.2), y + Inches(0.15), sw - Inches(0.4), Inches(2.2),
                     [(f"STEP {i+1}", {"bold": True, "size": Pt(12),
                                       "color": DesignSystem.COLOR_ACCENT}),
                      (st[:240], {"size": Pt(13), "space_before": Pt(6)})])
            animated.append((card.shape_id, AnimationType.APPEAR))
        if d.get("verify"):
            _textbox(slide, DesignSystem.MARGIN_LEFT, Inches(6.55),
                     DesignSystem.CONTENT_WIDTH, Inches(0.5),
                     [(f"✔ Verify: {d['verify']}"[:200], {"bold": True})], size=Pt(12),
                     color=DesignSystem.COLOR_SUCCESS)
        return animated

    @classmethod
    def render_matrix(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Comparison Matrix")
        headers, rows = d.get("headers", [])[:5], d.get("rows", [])[:6]
        if not headers:
            return []
        tbl = slide.shapes.add_table(len(rows) + 1, len(headers),
                                     DesignSystem.MARGIN_LEFT, Inches(2.1),
                                     DesignSystem.CONTENT_WIDTH, Inches(4.5)).table
        for j, h in enumerate(headers):
            cell = tbl.cell(0, j)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT
            for p in cell.text_frame.paragraphs:
                p.font.name = DesignSystem.FONT_BODY
                p.font.size = Pt(14)
                p.font.bold = True
                p.font.color.rgb = DesignSystem.COLOR_TEXT_ON_DARK
                p.alignment = PP_ALIGN.CENTER
        for i, row in enumerate(rows):
            for j in range(len(headers)):
                cell = tbl.cell(i + 1, j)
                cell.text = str(row[j]) if j < len(row) else ""
                cell.fill.solid()
                if j == 0:
                    cell.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT_BG
                else:
                    cell.fill.fore_color.rgb = (RGBColor(248, 250, 252)
                                                if i % 2 == 0 else RGBColor(255, 255, 255))
                for p in cell.text_frame.paragraphs:
                    p.font.name = DesignSystem.FONT_BODY
                    p.font.size = Pt(13)
                    p.font.bold = (j == 0)
                    p.font.color.rgb = DesignSystem.COLOR_TEXT_PRIMARY
        return []

    # -- assessment (locked answers only) ------------------------------------
    @classmethod
    def render_assessment(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        reveal = bool(d.get("reveal"))
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle,
                               "Answer & Solution" if reveal else "Question")
        origin = d.get("origin", "")
        _badge(slide, DesignSystem.MARGIN_LEFT, Inches(1.95), Inches(2.2),
               "Source question" if origin == "source" else "Retrieval check",
               DesignSystem.COLOR_ACCENT_BG, DesignSystem.COLOR_ACCENT)
        q_card = ShapeHelper.add_card(slide, DesignSystem.MARGIN_LEFT, Inches(2.35),
                                      DesignSystem.CONTENT_WIDTH, Inches(1.15),
                                      border_color=DesignSystem.COLOR_ACCENT, border_width=Pt(2))
        _textbox(slide, DesignSystem.MARGIN_LEFT + Inches(0.4), Inches(2.5),
                 DesignSystem.CONTENT_WIDTH - Inches(0.8), Inches(0.95),
                 [(d.get("prompt", "")[:400], {"bold": True})], size=Pt(17))
        options = d.get("options", {})
        animated: List[Tuple[int, AnimationType]] = []
        if options:
            letters = sorted(options)[:4]
            n = len(letters)
            if n <= 2:
                grid: List[List[str]] = [letters]
            else:
                grid = [letters[:2], letters[2:4]]
            oh, top, gap_y, gap_x = Inches(0.85), Inches(3.62), Inches(0.12), Inches(0.3)
            cw = (DesignSystem.CONTENT_WIDTH - gap_x) / 2
            for r, row in enumerate(grid):
                y = top + r * (oh + gap_y)
                for c, letter in enumerate(row):
                    x = DesignSystem.MARGIN_LEFT + c * (cw + gap_x) * (1 if len(row) > 1 else 0)
                    w = DesignSystem.CONTENT_WIDTH if len(row) == 1 else cw
                    is_right = reveal and letter == d.get("correct_option")
                    card = ShapeHelper.add_card(
                        slide, x, y, w, oh,
                        bg_color=DesignSystem.COLOR_SUCCESS_BG if is_right
                        else DesignSystem.COLOR_BG_CARD,
                        border_color=DesignSystem.COLOR_SUCCESS if is_right
                        else DesignSystem.COLOR_BORDER_CARD,
                        border_width=Pt(2) if is_right else Pt(1.5))
                    prefix = "✔ " if is_right else ""
                    _textbox(slide, x + Inches(0.25), y + Inches(0.08),
                             w - Inches(0.5), oh - Inches(0.16),
                             [(f"{prefix}{letter}) {options[letter]}"[:170],
                               {"bold": is_right})], size=Pt(13),
                             color=DesignSystem.COLOR_SUCCESS if is_right else None)
                    if is_right:
                        animated.append((card.shape_id, AnimationType.ANSWER_REVEAL))
            expl_top = top + len(grid) * (oh + gap_y) + Inches(0.1)
            if reveal and d.get("explanation"):
                _textbox(slide, DesignSystem.MARGIN_LEFT, expl_top,
                         DesignSystem.CONTENT_WIDTH, Inches(1.5),
                         [(d["explanation"][:340], {})], size=Pt(13),
                         color=DesignSystem.COLOR_TEXT_SECONDARY)
        else:
            _textbox(slide, DesignSystem.MARGIN_LEFT, Inches(4.2),
                     DesignSystem.CONTENT_WIDTH, Inches(1.6),
                     [("Quiet thinking time: 30 seconds. Write before we reveal." if not reveal
                       else d.get("explanation", "")[:340], {"bold": not reveal})],
                     size=Pt(16), color=DesignSystem.COLOR_ACCENT if not reveal else None)
        return animated

    # -- semantic prose structures -------------------------------------------
    @classmethod
    def render_practice(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Practice")
        origin = d.get("origin", "")
        _badge(slide, DesignSystem.MARGIN_LEFT, Inches(1.95), Inches(2.4),
               "Source problem" if origin == "source" else "Practice",
               DesignSystem.COLOR_WARNING_BG, DesignSystem.COLOR_WARNING)
        items = d.get("items", [])[:2]
        y = Inches(2.4)
        for item in items:
            card = ShapeHelper.add_card(slide, DesignSystem.MARGIN_LEFT, y,
                                        DesignSystem.CONTENT_WIDTH, Inches(1.9),
                                        border_color=DesignSystem.COLOR_WARNING_BORDER)
            _textbox(slide, DesignSystem.MARGIN_LEFT + Inches(0.4), y + Inches(0.2),
                     DesignSystem.CONTENT_WIDTH - Inches(0.8), Inches(1.55),
                     [(item[:420], {})], size=Pt(15))
            y += Inches(2.05)
        scaf = d.get("scaffold", {})
        if scaf.get("columns") and len(scaf["columns"]) >= 2:
            cols = scaf["columns"][:2]
            cw = (DesignSystem.CONTENT_WIDTH - Inches(0.4)) / 2
            for i, col in enumerate(cols):
                x = DesignSystem.MARGIN_LEFT + i * (cw + Inches(0.4))
                ShapeHelper.add_card(slide, x, y, cw, Inches(1.5))
                _textbox(slide, x + Inches(0.3), y + Inches(0.15), cw - Inches(0.6), Inches(1.2),
                         [(col, {"bold": True, "color": DesignSystem.COLOR_ACCENT}),
                          ("— class response —", {"size": Pt(13), "space_before": Pt(8)})])
            _textbox(slide, DesignSystem.MARGIN_LEFT, y + Inches(1.55),
                     DesignSystem.CONTENT_WIDTH, Inches(0.4),
                     [(scaf.get("note", "")[:160], {})], size=Pt(11),
                     color=DesignSystem.COLOR_TEXT_MUTED)
        elif scaf.get("note"):
            _textbox(slide, DesignSystem.MARGIN_LEFT, y, DesignSystem.CONTENT_WIDTH, Inches(0.5),
                     [(scaf["note"][:200], {})], size=Pt(12),
                     color=DesignSystem.COLOR_TEXT_SECONDARY)
        if d.get("faded"):
            _textbox(slide, DesignSystem.MARGIN_LEFT, Inches(6.8),
                     DesignSystem.CONTENT_WIDTH, Inches(0.4),
                     [("Guided: complete the missing step, then verify.", {"bold": True})],
                     size=Pt(12), color=DesignSystem.COLOR_ACCENT)
        return []

    @classmethod
    def render_definition(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Definition")
        term = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, DesignSystem.MARGIN_LEFT,
                                      Inches(2.1), DesignSystem.CONTENT_WIDTH, Inches(0.9))
        term.fill.solid()
        term.fill.fore_color.rgb = DesignSystem.COLOR_BG_HERO
        term.line.fill.background()
        p = term.text_frame.paragraphs[0]
        p.text = d.get("term", spec.title)[:120]
        p.font.name = DesignSystem.FONT_TITLE
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = DesignSystem.COLOR_TEXT_ON_DARK
        p.alignment = PP_ALIGN.CENTER
        card = ShapeHelper.add_card(slide, DesignSystem.MARGIN_LEFT, Inches(3.2),
                                    DesignSystem.CONTENT_WIDTH, Inches(3.2),
                                    border_color=DesignSystem.COLOR_ACCENT, border_width=Pt(2))
        _textbox(slide, DesignSystem.MARGIN_LEFT + Inches(0.5), Inches(3.5),
                 DesignSystem.CONTENT_WIDTH - Inches(1.0), Inches(2.7),
                 [(d.get("definition", "")[:600], {})], size=Pt(17))
        return [(term.shape_id, AnimationType.APPEAR)]

    @classmethod
    def render_points(cls, slide, spec: SlideSpec, badge: str = "Concept") -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, badge)
        points = (d.get("points") or d.get("prerequisites") or d.get("takeaways") or [])[:5]
        if not points:
            points = ["See speaker notes."]
        n = len(points)
        ch = min(Inches(3.6), (Inches(4.4) - Inches(0.25 * (n - 1))) / n)
        animated = []
        for i, pt in enumerate(points):
            y = Inches(2.1) + i * (ch + Inches(0.25))
            card = ShapeHelper.add_card(slide, DesignSystem.MARGIN_LEFT, y,
                                        DesignSystem.CONTENT_WIDTH, ch)
            _textbox(slide, DesignSystem.MARGIN_LEFT + Inches(0.7), y + Inches(0.12),
                     DesignSystem.CONTENT_WIDTH - Inches(1.0), ch - Inches(0.24),
                     [(f"{i+1:02d}", {"bold": True, "size": Pt(14),
                                      "color": DesignSystem.COLOR_ACCENT}),
                      (pt[:260], {"size": Pt(15), "space_before": Pt(2)})])
            animated.append((card.shape_id, AnimationType.APPEAR))
        return animated

    @classmethod
    def render_synthesis(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Synthesis")
        nodes = (d.get("nodes") or d.get("takeaways") or [])[:4]
        center = (d.get("center") or spec.title)[:60]
        cx, cy = Inches(5.7), Inches(3.6)
        hub = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(2.0), Inches(1.1))
        hub.fill.solid()
        hub.fill.fore_color.rgb = DesignSystem.COLOR_BG_HERO
        hub.line.fill.background()
        p = hub.text_frame.paragraphs[0]
        p.text = center
        p.font.name = DesignSystem.FONT_TITLE
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DesignSystem.COLOR_TEXT_ON_DARK
        p.alignment = PP_ALIGN.CENTER
        hub.text_frame.word_wrap = True
        animated = [(hub.shape_id, AnimationType.APPEAR)]
        spots = [(Inches(1.0), Inches(2.3)), (Inches(9.1), Inches(2.3)),
                 (Inches(1.0), Inches(5.2)), (Inches(9.1), Inches(5.2))]
        for (x, y), node in zip(spots, nodes):
            card = ShapeHelper.add_card(slide, x, y, Inches(3.2), Inches(1.3))
            _textbox(slide, x + Inches(0.25), y + Inches(0.12), Inches(2.7), Inches(1.1),
                     [(node[:160], {})], size=Pt(13))
            _connector(slide, x + Inches(1.6), y + Inches(0.65), cx + Inches(1.0), cy + Inches(0.55))
            animated.append((card.shape_id, AnimationType.APPEAR))
        return animated

    @classmethod
    def render_cause_chain(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        points = (d.get("points") or [])[:3]
        while len(points) < 3:
            points.append("—")
        labels = ("CAUSE", "MECHANISM", "EFFECT")
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Cause & Effect")
        animated = []
        pw = (DesignSystem.CONTENT_WIDTH - Inches(1.2)) / 3
        for i, (lab, pt) in enumerate(zip(labels, points)):
            x = DesignSystem.MARGIN_LEFT + i * (pw + Inches(0.6))
            card = ShapeHelper.add_card(slide, x, Inches(2.5), pw, Inches(3.4),
                                        border_color=DesignSystem.COLOR_ACCENT
                                        if i == 1 else DesignSystem.COLOR_BORDER_CARD,
                                        border_width=Pt(2) if i == 1 else Pt(1.5))
            _textbox(slide, x + Inches(0.25), Inches(2.7), pw - Inches(0.5), Inches(3.0),
                     [(lab, {"bold": True, "size": Pt(12), "color": DesignSystem.COLOR_ACCENT}),
                      (pt[:280], {"size": Pt(14), "space_before": Pt(6)})])
            animated.append((card.shape_id, AnimationType.APPEAR))
            if i < 2:
                arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + pw - Inches(0.02),
                                             Inches(4.0), Inches(0.64), Inches(0.4))
                arr.fill.solid()
                arr.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT
                arr.line.fill.background()
        return animated

    @classmethod
    def render_before_after(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        points = (d.get("points") or [])[:2]
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Before / After")
        cw = (DesignSystem.CONTENT_WIDTH - Inches(1.2)) / 2
        animated = []
        for i, lab in enumerate(("BEFORE", "AFTER")):
            x = DesignSystem.MARGIN_LEFT + i * (cw + Inches(1.2))
            card = ShapeHelper.add_card(slide, x, Inches(2.4), cw, Inches(3.8))
            _textbox(slide, x + Inches(0.3), Inches(2.6), cw - Inches(0.6), Inches(3.4),
                     [(lab, {"bold": True, "size": Pt(13), "color": DesignSystem.COLOR_ACCENT}),
                      ((points[i][:320] if i < len(points) else "—"),
                       {"size": Pt(15), "space_before": Pt(6)})])
            animated.append((card.shape_id, AnimationType.APPEAR))
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.35), Inches(4.0),
                                     Inches(0.64), Inches(0.4))
        arr.fill.solid()
        arr.fill.fore_color.rgb = DesignSystem.COLOR_ACCENT
        arr.line.fill.background()
        return animated

    @classmethod
    def render_hierarchy(cls, slide, spec: SlideSpec) -> List[Tuple[int, AnimationType]]:
        d = spec.elements_data
        points = (d.get("points") or [])[:4]
        ShapeHelper.set_slide_background(slide)
        ShapeHelper.add_header(slide, spec.title, spec.subtitle, "Hierarchy")
        animated = []
        for i, pt in enumerate(points):
            indent = Inches(0.6 * i)
            w = DesignSystem.CONTENT_WIDTH - indent * 2
            x = DesignSystem.MARGIN_LEFT + indent
            y = Inches(2.2) + i * Inches(1.2)
            card = ShapeHelper.add_card(slide, x, y, w, Inches(1.0))
            _textbox(slide, x + Inches(0.3), y + Inches(0.15), w - Inches(0.6), Inches(0.7),
                     [(f"LEVEL {i+1}  •  {pt[:160]}", {})], size=Pt(14))
            animated.append((card.shape_id, AnimationType.APPEAR))
        return animated
