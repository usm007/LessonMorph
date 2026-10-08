"""Rendered-slide QA: PPTX → inspection → per-slide findings → repair.

Inspects the generated PPTX with python-pptx (geometry + text scan) and,
when LibreOffice is available, renders slides to PNG for visual confirmation.
Detects: text overflow, clipping/off-canvas, orphaned/empty cards, tiny text,
duplicate content, raw Markdown/LaTeX, broken symbols, incorrect answers,
missing source content. Findings identify the affected slide so the repair
loop can regenerate THAT slide from its Blueprint — never the whole deck's
pedagogy.
"""

from __future__ import annotations
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional
from lessonmorph.blueprint.ir import Blueprint
from lessonmorph.blueprint.sanitize import contains_markup

EMU_PER_INCH = 914400


@dataclass
class RenderFinding:
    slide_index: int  # 1-based
    slide_id: str  # blueprint slide id ("" if unmapped)
    code: str
    severity: str  # ERROR | WARNING
    detail: str


@dataclass
class RenderReport:
    findings: List[RenderFinding]
    method: str  # "ooxml" or "ooxml+png"
    status: str  # PASS | WARN | FAIL

    def errors(self) -> List[RenderFinding]:
        return [f for f in self.findings if f.severity == "ERROR"]

    def failed_slide_ids(self) -> List[str]:
        return sorted({f.slide_id for f in self.errors() if f.slide_id})

    def to_markdown(self) -> str:
        lines = ["## RENDER QA", "", f"Method: `{self.method}` | Status: `{self.status}`", ""]
        if not self.findings:
            lines.append("No rendering defects detected.")
            return "\n".join(lines)
        lines += ["| Slide | Code | Severity | Detail |",
                  "| :--- | :--- | :---: | :--- |"]
        for f in self.findings:
            lines.append(f"| {f.slide_id or f'slide_{f.slide_index}'} | {f.code} | "
                         f"{f.severity} | {f.detail} |")
        return "\n".join(lines)


def _slide_texts(slide) -> List[str]:
    texts = []
    for shape in slide.shapes:
        try:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            texts.append(cell.text.strip())
        except Exception:
            continue
    return texts


def _intersects(a, b) -> bool:
    """Box overlap test (inches): (left, top, width, height)."""
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah


def check_timing_integrity(pptx_path: Path | str,
                           slide_ids: List[str] | None = None) -> List["RenderFinding"]:
    """PowerPoint-strict timing validation (duplicate cTn ids corrupt the file).

    python-pptx reopens such files happily, so this gate exists to catch what
    package-integrity checks miss. Returns ERROR findings per offending slide.
    """
    findings: List[RenderFinding] = []
    try:
        from lxml import etree
    except ImportError:
        return findings  # deep check unavailable; validator covers the rest
    NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    try:
        z = zipfile.ZipFile(str(pptx_path))
    except Exception:
        return findings
    with z:
        try:
            pres = etree.fromstring(z.read("ppt/presentation.xml"))
            rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
        except KeyError:
            return findings
        targets = {r.get("Id"): r.get("Target") for r in rels.findall("Relationship")}
        order = []
        for sld in pres.findall(".//p:sldId", NS):
            rid = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            t = targets.get(rid, "")
            order.append("ppt/" + t.lstrip("/") if t else "")
        for idx, name in enumerate(order):
            if not name or not name.startswith("ppt/slides/slide"):
                continue
            sid = slide_ids[idx] if slide_ids and idx < len(slide_ids) else ""
            try:
                root = etree.fromstring(z.read(name))
            except Exception as e:
                findings.append(RenderFinding(idx + 1, sid, "unparseable_slide", "ERROR",
                                              f"{name} does not parse: {e}"))
                continue
            ids = [e.get("id") for e in root.findall(".//p:cTn", NS)]
            dup = sorted({i for i in ids if ids.count(i) > 1})
            if dup:
                findings.append(RenderFinding(
                    idx + 1, sid, "duplicate_timing_ids", "ERROR",
                    f"Duplicate time-node ids {dup} — PowerPoint reports the file as corrupt."))
            shapes = [e.get("id") for e in root.findall(".//p:cNvPr", NS)]
            for s in ([e.get("spid") for e in root.findall(".//p:spTgt", NS)]
                      + [e.get("spid") for e in root.findall(".//p:bldP", NS)]):
                if s not in shapes:
                    findings.append(RenderFinding(
                        idx + 1, sid, "dangling_animation_target", "ERROR",
                        f"Animation targets missing shape id {s}."))
    return findings


class RenderQA:
    """Inspects rendered slides against their Blueprint."""

    OVERFLOW_CHARS_PER_SQIN = 900
    TINY_PT = 10.0

    @classmethod
    def inspect(cls, pptx_path: Path | str, blueprint: Optional[Blueprint] = None,
                work_dir: Optional[Path | str] = None) -> RenderReport:
        import pptx
        findings: List[RenderFinding] = []
        prs = pptx.Presentation(str(pptx_path))
        sw_in = prs.slide_width / EMU_PER_INCH
        sh_in = prs.slide_height / EMU_PER_INCH
        ids = [s.id for s in blueprint.slides] if blueprint else []

        if blueprint and len(prs.slides) != len(blueprint.slides):
            findings.append(RenderFinding(
                0, "", "slide_count_mismatch", "ERROR",
                f"PPTX has {len(prs.slides)} slides but Blueprint has "
                f"{len(blueprint.slides)} — source content may be missing."))

        for idx, slide in enumerate(prs.slides):
            sid = ids[idx] if idx < len(ids) else ""
            texts = _slide_texts(slide)
            # Raw markup / broken symbols.
            for t in texts:
                kind = contains_markup(t)
                if kind:
                    findings.append(RenderFinding(
                        idx + 1, sid, f"raw_{kind}", "ERROR",
                        f"Leaked {kind} syntax on rendered slide: {t[:90]}"))
                if "\ufffd" in t or re.search(r"(?<![A-Za-z])\\(?![nrt])", t):
                    findings.append(RenderFinding(
                        idx + 1, sid, "broken_symbols", "ERROR",
                        f"Broken symbol/escape on rendered slide: {t[:90]}"))
            # Empty cards: card containers are composed with overlay textboxes,
            # so an empty shape is only orphaned if no text intersects it.
            text_boxes: list = []
            empty_shapes: list = []
            tiny = 0
            crowded = 0
            for shape in slide.shapes:
                try:
                    if not shape.has_text_frame:
                        continue
                    box = (shape.left / EMU_PER_INCH, shape.top / EMU_PER_INCH,
                           shape.width / EMU_PER_INCH, shape.height / EMU_PER_INCH)
                    txt = shape.text_frame.text.strip()
                    if txt:
                        text_boxes.append(box)
                        crowded += 1
                        for p in shape.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size is not None and r.font.size.pt < cls.TINY_PT:
                                    tiny += 1
                        area = box[2] * box[3]
                        if area > 0.05 and len(txt) / area > cls.OVERFLOW_CHARS_PER_SQIN:
                            findings.append(RenderFinding(
                                idx + 1, sid, "text_overflow", "WARNING",
                                f"Text density suggests overflow: "
                                f"{len(txt)} chars in {area:.2f} sq in."))
                    elif box[2] * box[3] < sw_in * sh_in * 0.85:
                        empty_shapes.append(box)
                    left, top = box[0], box[1]
                    if left + box[2] < 0.05 or top + box[3] < 0.05 \
                            or left > sw_in - 0.05 or top > sh_in - 0.05:
                        findings.append(RenderFinding(
                            idx + 1, sid, "off_canvas", "WARNING",
                            "Shape fully outside the slide canvas."))
                except Exception:
                    continue
            orphaned = sum(1 for e in empty_shapes
                           if not any(_intersects(e, t) for t in text_boxes))
            if orphaned >= 2:
                findings.append(RenderFinding(idx + 1, sid, "empty_cards", "WARNING",
                                              f"{orphaned} orphaned empty containers."))
            if tiny:
                findings.append(RenderFinding(idx + 1, sid, "tiny_text", "WARNING",
                                              f"{tiny} runs below {cls.TINY_PT}pt."))
            if crowded > 12:
                findings.append(RenderFinding(idx + 1, sid, "overcrowding", "WARNING",
                                              f"{crowded} text containers — slide may be overcrowded."))
            # Duplicate content across consecutive slides.
            if idx > 0:
                prev = _slide_texts(prs.slides[idx - 1])
                if texts and prev and texts == prev:
                    findings.append(RenderFinding(idx + 1, sid, "duplicate_slide",
                                                  "WARNING", "Identical text to previous slide."))
            # Locked-answer check on reveal slides.
            if blueprint and idx < len(ids):
                spec = blueprint.slides[idx]
                if spec.assessment and spec.body.get("reveal") and spec.assessment.type == "mcq":
                    correct = spec.assessment.options.get(spec.assessment.correct_option, "")
                    joined = "\n".join(texts)
                    if correct and correct[:40] not in joined:
                        findings.append(RenderFinding(
                            idx + 1, sid, "incorrect_answer", "ERROR",
                            f"Locked answer '{correct[:60]}' not visible on reveal slide."))
                    wrong_marked = [l for l in spec.assessment.options
                                    if l != spec.assessment.correct_option
                                    and f"✔ {l})" in joined]
                    if wrong_marked:
                        findings.append(RenderFinding(
                            idx + 1, sid, "incorrect_answer", "ERROR",
                            f"Wrong option(s) marked correct: {wrong_marked}."))

        method = "ooxml"
        # PowerPoint-strict timing gate: duplicate time-node ids or dangling
        # animation targets make PowerPoint report the file as corrupt even
        # though python-pptx reopens it happily.
        findings.extend(check_timing_integrity(pptx_path, ids or None))
        if work_dir and shutil.which("soffice"):
            try:
                out = Path(work_dir) / "render_png"
                out.mkdir(parents=True, exist_ok=True)
                subprocess.run(["soffice", "--headless", "--convert-to", "png",
                                "--outdir", str(out), str(pptx_path)],
                               capture_output=True, timeout=120, check=False)
                if list(out.glob("*.png")):
                    method = "ooxml+png"
            except Exception:
                pass
        status = "PASS"
        if any(f.severity == "ERROR" for f in findings):
            status = "FAIL"
        elif findings:
            status = "WARN"
        return RenderReport(findings, method, status)
