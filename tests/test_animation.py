"""Tests for OpenXML animation injection."""

from pathlib import Path
import pptx
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from lessonmorph.animation.engine import OoxmlAnimationEngine
from lessonmorph.core.models import AnimationType


def test_animation_injection(tmp_path: Path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    
    # Apply animation
    OoxmlAnimationEngine.apply_animations(slide, [(shape.shape_id, AnimationType.APPEAR)])
    
    out_file = tmp_path / "animated.pptx"
    prs.save(str(out_file))

    # Re-open and inspect
    prs_reload = pptx.Presentation(out_file)
    timing_tags = [child for child in prs_reload.slides[0]._element if "timing" in child.tag]
    assert len(timing_tags) == 1
