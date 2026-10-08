"""Tests for the Presentation Blueprint IR layer (execution plan, not pedagogy)."""

import io
import json

import pytest

from lessonmorph.blueprint.adapter import (
    BlueprintCompiler,
    classify_task,
    lock_assessment,
    parse_mcq_options,
    select_representation,
    structure_equation,
)
from lessonmorph.blueprint.composer import SlideComposer
from lessonmorph.blueprint.ir import (
    TASK_FAMILIES,
    Blueprint,
    InventoryItem,
    LockedAssessment,
    Representation,
    SlideIR,
    TaskType,
    VisualSpec,
)
from lessonmorph.blueprint.registry import (
    REGISTRY,
    fallback_for,
    representations_for,
    tasks_for,
)
from lessonmorph.blueprint.repair import repair_blueprint
from lessonmorph.blueprint.render_qa import RenderQA
from lessonmorph.blueprint.sanitize import clean, contains_markup, latex_to_unicode
from lessonmorph.blueprint.validators import BlueprintValidator
from lessonmorph.core.models import (
    ChapterPlan,
    ContentImportance,
    ContentType,
    DocumentSection,
    QuestionType,
    QuizQuestion,
    SubjectDomain,
)
from lessonmorph.ledger.ledger import ContentCompletenessLedger
from lessonmorph.pedagogy.planner import PedagogicalPlanner


def _unit(kind=ContentType.EXPLANATION, text="Photosynthesis converts light energy.",
          uid=None, importance=ContentImportance.CORE):
    from lessonmorph.core.models import ContentUnit
    return ContentUnit(id=uid or "C001", source_location="p1", chapter_id="ch01",
                       chapter_title="T", section_title=None, content_type=kind,
                       importance=importance, normalized_content=text,
                       original_wording=text)


# -- sanitize ---------------------------------------------------------------

def test_latex_to_unicode_never_emits_backslashes():
    out = latex_to_unicode("$$6\\text{CO}_2 + 6\\text{H}_2\\text{O} \\longrightarrow "
                           "\\text{C}_6\\text{H}_{12}\\text{O}_6 + 6\\text{O}_2$$")
    assert "\\" not in out
    assert "CO₂" in out and "H₂O" in out and "→" in out


def test_frac_with_nested_text_resolves():
    out = latex_to_unicode("\\frac{8.0\\text{ mol photons}}{8\\text{ photons/molecule}}")
    assert "frac" not in out
    assert "8.0 mol photons/8 photons/molecule" == out


def test_clean_strips_markdown_and_single_dollars():
    assert clean("**Definition**: $CO_2$ here") == "Definition: CO₂ here"
    assert contains_markup("a $$b$$ c") == "latex"
    assert contains_markup("**bold** text") == "markdown"
    assert contains_markup("| a | b |\n|---|---|") == "markdown"
    assert contains_markup("plain clean text") == ""


# -- task taxonomy + representation legality ----------------------------------

def test_task_taxonomy_drives_structure():
    assert classify_task(_unit(text="Trace the electron from PSII to NADPH.")).value == "process"
    assert classify_task(_unit(text="Contrast C3 and C4 in arid zones.")).value == "comparison"
    u = _unit(kind=ContentType.EXERCISE,
              text="Question: What happens?\nA) x\nB) y")
    assert classify_task(u).value == "prediction"


def test_representation_stays_inside_task_family():
    u = _unit(text="Trace the electron from PSII water photolysis to NADPH.")
    task = classify_task(u)
    rep = select_representation(task, u)
    assert rep.value in [r.value for r in TASK_FAMILIES[task]]
    assert rep != Representation.WORKED_CALCULATION  # never numerical for process


def test_definition_and_table_selection():
    d = _unit(kind=ContentType.DEFINITION, text="Photosynthesis is an anabolic process.")
    rep = select_representation(classify_task(d), d)
    assert rep == Representation.DEFINITION_FOCUS


# -- equations ------------------------------------------------------------------

def test_structure_equation_splits_terms_and_energy():
    eq = structure_equation("$$6\\text{CO}_2 + 6\\text{H}_2\\text{O} + \\text{light energy} "
                            "\\longrightarrow \\text{C}_6\\text{H}_{12}\\text{O}_6 + 6\\text{O}_2$$",
                            ["CO₂ represents carbon dioxide"])
    assert [t.species for t in eq.lhs] == ["CO₂", "H₂O"]
    assert [t.coefficient for t in eq.lhs] == ["6", "6"]
    assert eq.energy == "light energy"
    assert eq.rhs[0].species == "C₆H₁₂O₆"
    assert "→" in eq.inline() and "\\" not in eq.inline()


# -- locked assessments ------------------------------------------------------------

def test_mcq_parsing_and_lock_from_key():
    text = "Question: What happens?\nA) water stops\nB) ATP stops, flow continues"
    stem, options = parse_mcq_options(text)
    assert stem == "What happens?"
    assert options == {"A": "water stops", "B": "ATP stops, flow continues"}
    q = QuizQuestion("Q01", QuestionType.MULTIPLE_CHOICE, ["C009"], "application",
                     prompt=text, options=["A) water stops", "B) ATP stops, flow continues"],
                     correct_answer="A) water stops", explanation="Fallback guess.")
    locked = lock_assessment(q, _unit(kind=ContentType.EXERCISE, text=text, uid="C009"),
                             [{"match": "what happens", "correct_option": "B",
                               "explanation": "Gradient collapses.", "source_anchor": "S3"}])
    assert locked.is_locked() and locked.correct_option == "B"
    assert locked.origin == "source"


def test_unresolvable_answer_stays_unlocked_for_qa():
    q = QuizQuestion("Q09", QuestionType.MULTIPLE_CHOICE, ["C009"], "application",
                     prompt="Q?\nA) x\nB) y", options=["A) x", "B) y"],
                     correct_answer="something unrelated", explanation="e")
    locked = lock_assessment(q, _unit(kind=ContentType.EXERCISE, text="Q?\nA) x\nB) y",
                                      uid="C009"), [])
    assert not locked.is_locked()


# -- blueprint validation -------------------------------------------------------

def _bad_blueprint() -> Blueprint:
    bp = Blueprint(chapter_id="ch01", chapter_title="T")
    bp.inventory.append(InventoryItem(source_id="C001", kind="definition"))
    bp.slides.append(SlideIR(id="slide_01", purpose="p", task="concept",
                             representation="", content_title="T"))  # no representation
    return bp


def test_blueprint_qa_rejects_missing_representation_and_unpreserved():
    report = BlueprintValidator.validate(_bad_blueprint())
    assert report.status == "FAIL"
    codes = {f.code for f in report.errors()}
    assert "missing_representation" in codes
    assert "unpreserved_source_element" in codes


def test_blueprint_qa_rejects_unlocked_mcq():
    bp = Blueprint(chapter_id="ch01", chapter_title="T")
    bp.inventory.append(InventoryItem(source_id="C009", kind="question",
                                      slide_id="slide_01", preserved=True))
    bp.slides.append(SlideIR(
        id="slide_01", purpose="assess", task="prediction", representation="mcq",
        content_title="Q", body={"prompt": "Q?", "options": {"A": "x", "B": "y"}},
        source_ids=["C009"],
        assessment=LockedAssessment(id="q", type="mcq", stem="Q?",
                                    options={"A": "x", "B": "y"},
                                    correct_option="", explanation="e",
                                    origin="source")))
    report = BlueprintValidator.validate(bp)
    assert any(f.code == "answer_not_locked" and f.severity == "ERROR"
               for f in report.findings)


# -- composer determinism ----------------------------------------------------------

def test_composer_is_deterministic_and_sanitized():
    ir = SlideIR(id="slide_01", purpose="teach", concept_title="T",
                 task="definition", representation="definition_focus",
                 content_title="T", body={"term": "X", "definition": "**Bold** $x$"})
    c = SlideComposer()
    a = c.compose(ir, "ch01")
    b = SlideComposer().compose(ir, "ch01")
    assert a.elements_data == b.elements_data
    assert a.slide_id == b.slide_id == "S001"
    assert a.visual_model == "definition_focus"
    assert "**" not in a.elements_data["definition"]


# -- repair --------------------------------------------------------------------------

def test_repair_falls_back_representation_only():
    bp = Blueprint(chapter_id="ch01", chapter_title="T")
    bp.slides.append(SlideIR(id="slide_01", purpose="p", task="process",
                             representation="pathway", content_title="P",
                             body={}, visual=VisualSpec(structure=["a", "b"])))
    from lessonmorph.blueprint.render_qa import RenderFinding, RenderReport
    report = RenderReport(
        [RenderFinding(1, "slide_01", "text_overflow", "ERROR", "dense")],
        "ooxml", "FAIL")
    repaired, notes = repair_blueprint(bp, report)
    assert repaired == ["slide_01"]
    assert bp.slides[0].representation == "sequence"  # same family, simpler


# -- adapter end-to-end (synthetic pedagogy output) -----------------------------------

def _photo_like_ledger() -> ContentCompletenessLedger:
    ledger = ContentCompletenessLedger(document_title="Photo")
    ledger.add_unit("p1", "ch01", "Photo", ContentType.DEFINITION,
                    ContentImportance.CORE, "Photosynthesis is an anabolic process.")
    ledger.add_unit("p1", "ch01", "Photo", ContentType.FORMULA,
                    ContentImportance.CORE,
                    "$$6\\text{CO}_2 \\longrightarrow \\text{C}_6\\text{H}_{12}\\text{O}_6$$")
    return ledger


def test_adapter_compiles_pedagogy_to_executable_blueprint(tmp_path):
    ledger = _photo_like_ledger()
    plan = PedagogicalPlanner(ledger).plan_chapter(DocumentSection("ch01", "Photo", 1, 1))
    bp = BlueprintCompiler(ledger).compile(plan)
    assert bp.slides
    assert all(s.representation for s in bp.slides)  # hard prohibition enforced
    report = BlueprintValidator.validate(bp)
    assert report.status in ("PASS", "WARN")
    assert not report.errors()
    kinds = {i.kind for i in bp.inventory}
    assert "definition" in kinds and "equation" in kinds
    assert all(i.preserved for i in bp.inventory if i.kind in ("definition", "equation"))


def test_grammar_registry_covers_all_representations():
    for rep in Representation:
        assert rep.value in REGISTRY
        assert tasks_for(rep) or rep.value in ("title", "roadmap", "objectives", "exit")
    for task in TaskType:
        assert representations_for(task), f"{task} has no legal representation"
    assert fallback_for(Representation.PATHWAY) == Representation.SEQUENCE


def test_animation_timing_ids_are_unique_and_resolve(tmp_path):
    """PowerPoint reports files with duplicate time-node ids as corrupt."""
    from pptx import Presentation
    from pptx.util import Inches
    from lessonmorph.animation.engine import OoxmlAnimationEngine
    from lessonmorph.blueprint.render_qa import check_timing_integrity
    from lessonmorph.core.models import AnimationType
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Hello"
    OoxmlAnimationEngine.apply_animations(
        slide, [(box.shape_id, AnimationType.APPEAR),
                (box.shape_id, AnimationType.HIGHLIGHT)])
    out = tmp_path / "anim.pptx"
    prs.save(str(out))
    assert check_timing_integrity(out) == []


def test_render_qa_detects_leaked_markup(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tb.text_frame.text = "Solve $$x^2$$ now"
    out = tmp_path / "leak.pptx"
    prs.save(str(out))
    report = RenderQA.inspect(out)
    assert any(f.code == "raw_latex" and f.severity == "ERROR" for f in report.findings)
