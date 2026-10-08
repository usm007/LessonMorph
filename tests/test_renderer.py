"""Tests for PPTX renderer and native slide generation."""

from pathlib import Path
import pptx
from pptx.util import Inches
from lessonmorph.core.models import SlideSpec, SlideType, SpeakerNotes
from lessonmorph.renderer.engine import PptxRenderer


def test_pptx_rendering_and_notes(tmp_path: Path):
    renderer = PptxRenderer()
    slides = [
        SlideSpec(
            slide_id="S001",
            chapter_id="ch01",
            title="Introduction to Photosynthesis",
            subtitle="Biology • Classroom Presentation",
            purpose="Introduce topic",
            source_content_ids=[],
            slide_type=SlideType.TITLE,
            visual_model="hero_title",
            elements_data={"unit_title": "BIOLOGY", "topic": "Photosynthesis"},
            speaker_notes=SpeakerNotes(
                teacher_explanation="Welcome students to photosynthesis.",
                emphasis="Focus on chloroplast structure.",
                ask_students="Where does the plant get its mass?",
            ),
        ),
        SlideSpec(
            slide_id="S002",
            chapter_id="ch01",
            title="Definition of Photosynthesis",
            subtitle="Core process",
            purpose="Define term",
            source_content_ids=["C001"],
            slide_type=SlideType.CONCEPT_DEFINITION,
            visual_model="definition_card",
            elements_data={
                "definition_text": "Photosynthesis is the process by which light energy is converted to chemical energy.",
                "source_ref": "Page 12",
            },
            speaker_notes=SpeakerNotes(
                teacher_explanation="Emphasize that glucose is synthesized.",
            ),
        ),
    ]

    out_file = tmp_path / "test_deck.pptx"
    renderer.render_presentation(slides, out_file)
    assert out_file.exists()

    # Re-open and verify with python-pptx
    prs = pptx.Presentation(out_file)
    assert len(prs.slides) == 2
    assert prs.slide_width == Inches(13.333)
    assert prs.slide_height == Inches(7.5)

    # Verify speaker notes
    slide1 = prs.slides[0]
    assert slide1.has_notes_slide
    notes_text = slide1.notes_slide.notes_text_frame.text
    assert "TEACHER EXPLANATION:" in notes_text
    assert "Where does the plant get its mass?" in notes_text
