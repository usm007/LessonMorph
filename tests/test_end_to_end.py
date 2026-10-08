"""End-to-End integration tests for LessonMorph compiler."""

from pathlib import Path
import pptx
from lessonmorph.cli import compile_document


SAMPLE_LESSON_MD = """# Newton's Laws of Motion

## Chapter 1: The Foundations of Dynamics

**Definition**: Inertia is the inherent property of matter by which it continues in its existing state of rest or uniform motion in a straight line, unless that state is changed by an external force.

Sir Isaac Newton formulated three fundamental laws of motion that govern classical mechanics.

### The Second Law of Motion

The rate of change of momentum of a body is directly proportional to the applied force and takes place in the direction in which the force acts.

Formula:
$$F = m \\cdot a$$

Where:
- $F$ is the net force measured in Newtons (N)
- $m$ is the inertial mass of the object measured in kilograms (kg)
- $a$ is the resulting acceleration measured in meters per second squared ($m/s^2$)

### Worked Example: Accelerating a Sled

A sled with a mass of 25 kg is pulled across friction-free ice. The applied net forward force is 100 N. Calculate the resulting acceleration of the sled.

Step 1: Identify given quantities: Mass $m = 25\\text{ kg}$, Net force $F = 100\\text{ N}$.
Step 2: State Newton's Second Law: $F = m \\cdot a$, so $a = F / m$.
Step 3: Substitute the known values: $a = 100\\text{ N} / 25\\text{ kg} = 4.0\\text{ m/s}^2$.

### Common Misconception: Force and Velocity

Common Mistake: A constant forward velocity requires a constant forward net force.
Correction: A body at constant velocity experiences zero net force ($\\Sigma F = 0$). Net force causes acceleration, not velocity.

### Data Summary: Acceleration Under Constant Force

| Mass (kg) | Force (N) | Acceleration (m/s²) |
| :---: | :---: | :---: |
| 10 | 50 | 5.0 |
| 20 | 50 | 2.5 |
| 25 | 50 | 2.0 |
| 50 | 50 | 1.0 |

### Check for Understanding

Question: A spacecraft travels at constant speed in deep interstellar space far from any gravitational source. What is the net external force acting on it?
A) Equal to its mass times speed
B) Exactly zero Newtons
C) Directed towards the nearest star
D) Constantly decreasing

### Practice Problem

Practice Problem 1: A 1200 kg car accelerates from rest at a rate of 3.5 m/s². Determine the net force exerted on the car.
"""


def test_full_compiler_pipeline(tmp_path: Path):
    src_file = tmp_path / "newtons_laws.md"
    src_file.write_text(SAMPLE_LESSON_MD, encoding="utf-8")

    out_pptx = tmp_path / "output" / "newtons_laws.pptx"
    work_dir = tmp_path / "work" / "newtons_laws"
    lesson_dir = tmp_path / "lesson"

    res = compile_document(src_file, output_file=out_pptx, work_dir=work_dir,
                           lesson_dir=lesson_dir)

    assert res["status"] in ("PASS", "WARN")
    assert out_pptx.exists()
    assert res["total_units"] > 5
    assert res["uncovered_units"] == 0  # 100% Coverage

    # Verify PRIMARY browser bundle (LessonQA-gated)
    assert res["lesson_status"] == "PASS"
    assert (lesson_dir / "index.html").exists()
    assert (lesson_dir / "lesson.json").exists()
    assert (lesson_dir / "teacher.json").exists()

    # Verify PPTX integrity
    prs = pptx.Presentation(out_pptx)
    assert len(prs.slides) >= 10

    # Verify that notes exist across slides
    notes_count = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    assert notes_count >= len(prs.slides) * 0.8

    # Verify validation report file exists
    assert Path(res["validation_report"]).exists()
