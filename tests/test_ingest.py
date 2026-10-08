"""Tests for PDF, DOCX, and Text ingestion."""

from pathlib import Path
import docx
import pymupdf
from lessonmorph.ingest.detector import ingest_document


def test_pdf_ingestion(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Chapter 1: Principles of Genetics\n\nDefinition: A gene is a basic unit of heredity.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "DNA replication occurs in the cell nucleus.")
    doc.save(str(pdf_path))
    doc.close()

    res = ingest_document(pdf_path)
    assert res.total_pages == 2
    assert len(res.doc_map.sections) >= 1
    assert "Genetics" in res.doc_map.sections[0].text_content or "Genetics" in res.doc_map.sections[0].title


def test_docx_ingestion(tmp_path: Path):
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Cell Biology", level=1)
    doc.add_paragraph("All living organisms are made of cells.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Organelle"
    table.rows[0].cells[1].text = "Function"
    table.rows[1].cells[0].text = "Mitochondria"
    table.rows[1].cells[1].text = "Energy"
    doc.save(str(docx_path))

    res = ingest_document(docx_path)
    assert len(res.doc_map.sections) >= 1
    assert len(res.tables) == 1
    assert res.tables[0].headers == ["Organelle", "Function"]
