"""Multi-chapter + new visual-model regression tests."""
from pathlib import Path
import pptx
from lessonmorph.cli import compile_document
from lessonmorph.ledger.ledger import ContentCompletenessLedger
from lessonmorph.pedagogy.planner import PedagogicalPlanner
from lessonmorph.storyboard.engine import StoryboardEngine
from lessonmorph.renderer.engine import PptxRenderer
from lessonmorph.ingest.detector import ingest_document


MULTI_MD = """# General Science

## Chapter 1: Motion Basics

**Definition**: Speed is the rate of change of distance with time.

Motion can be classified into linear and circular types.

## Chapter 2: Energy Basics

**Definition**: Energy is the capacity to do work.

Energy flows from producers to consumers in stages.
"""


def test_multi_chapter_deck_has_dividers_and_full_coverage(tmp_path: Path):
    src = tmp_path / "multi.md"
    src.write_text(MULTI_MD, encoding="utf-8")
    out = tmp_path / "multi.pptx"
    work = tmp_path / "work"
    res = compile_document(src, out, work, lesson_dir=tmp_path / "lesson")
    assert res["chapters_count"] == 2
    assert res["uncovered_units"] == 0
    prs = pptx.Presentation(out)
    titles = []
    for s in prs.slides:
        texts = [sh.text for sh in s.shapes if sh.has_text_frame]
        titles.extend(texts)
    blob = "\n".join(titles)
    assert "Contents" in blob
    assert "Motion Basics" in blob and "Energy Basics" in blob


def test_subject_visual_models_render_and_reopen(tmp_path: Path):
    src = tmp_path / "models.md"
    src.write_text("# Title\n\n## Ch\n\nPhotosynthesis occurs in stages across cell membranes classified into types.", encoding="utf-8")
    ingest_res = ingest_document(src)
    ledger = ContentCompletenessLedger(document_title="t")
    from lessonmorph.ingest.atomizer import ContentAtomizer
    ContentAtomizer(ledger).atomize(ingest_res)
    plan = PedagogicalPlanner(ledger).plan_chapter(ingest_res.doc_map.sections[0])
    slides = StoryboardEngine(ledger).generate_storyboard(plan)
    # Force each new visual model at least once
    for i, model in enumerate(["diagram_explanation", "process_flow", "2_column_compare", "classification_grid", "timeline", "cause_effect"]):
        if i < len(slides):
            slides[i].visual_model = model
            slides[i].elements_data = {"content": "Alpha stage. Beta stage.", "caption": "Caption", "entries": ["A", "B"]}
    out = tmp_path / "models.pptx"
    PptxRenderer().render_presentation(slides, out)
    prs = pptx.Presentation(out)
    assert len(prs.slides) == len(slides)
